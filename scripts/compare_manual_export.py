"""Compare TG Booster manual Excel export vs PostgreSQL (one-off reconciliation helper)."""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from etl.config import get_settings

import psycopg


def load_excel_daily(path: Path) -> dict[str, dict[str, float | int]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    idx = {h: i for i, h in enumerate(rows[0])}
    by_day: dict[str, dict[str, float | int]] = defaultdict(
        lambda: {"spend": 0.0, "leads": 0, "views": 0, "clicks": 0}
    )
    for row in rows[1:]:
        if not row or row[idx["date_day"]] is None:
            continue
        day = str(row[idx["date_day"]])
        if day == "ВСЕГО":
            continue
        day = day[:10]
        by_day[day]["spend"] += float(row[idx["metric_spent"]] or 0)
        by_day[day]["leads"] += int(row[idx["metric_joins"]] or 0)
        by_day[day]["views"] += int(row[idx["metric_views"]] or 0)
        by_day[day]["clicks"] += int(row[idx["metric_clicks"]] or 0)
    return dict(by_day)


def load_pg_daily(
    database_url: str,
    account_id: str,
    date_from: date,
    date_to: date,
) -> dict[str, tuple[float, int]]:
    sql = """
        SELECT stat_date::text, SUM(spend), SUM(leads)
        FROM raw_campaign_stats
        WHERE account_id = %s AND stat_date BETWEEN %s AND %s
        GROUP BY stat_date
        ORDER BY stat_date
    """
    with psycopg.connect(database_url) as conn:
        with conn.cursor() as cur:
            cur.execute(sql, (account_id, date_from, date_to))
            rows = cur.fetchall()
    return {r[0]: (float(r[1]), int(r[2])) for r in rows}


def main() -> None:
    if len(sys.argv) < 3:
        print("Usage: python scripts/compare_manual_export.py <excel> <account_id>")
        sys.exit(1)

    excel_path = Path(sys.argv[1])
    account_id = sys.argv[2]
    manual = load_excel_daily(excel_path)
    days = sorted(manual)
    d_from = date.fromisoformat(days[0])
    d_to = date.fromisoformat(days[-1])
    settings = get_settings()

    pg = load_pg_daily(settings.database_url, account_id, d_from, d_to)

    print(f"Period: {d_from} -> {d_to}  account_id={account_id}")
    print("date,manual_spend,pg_spend,delta_spend_pct,manual_leads,pg_leads,delta_leads_pct")
    t_ms = t_ps = t_ml_i = t_pl_i = 0.0
    for day in days:
        m = manual[day]
        p = pg.get(day)
        if not p:
            continue
        ds = p[0] - float(m["spend"])
        dl = p[1] - int(m["leads"])
        dsp = 100 * ds / float(m["spend"]) if m["spend"] else 0
        dlp = 100 * dl / int(m["leads"]) if m["leads"] else 0
        print(
            f"{day},{m['spend']:.2f},{p[0]:.2f},{dsp:+.3f},"
            f"{m['leads']},{p[1]},{dlp:+.3f}"
        )
        t_ms += float(m["spend"])
        t_ps += p[0]
        t_ml_i += int(m["leads"])
        t_pl_i += p[1]

    print(
        f"# manual_total spend={t_ms:.2f} leads={t_ml_i} | "
        f"pg_sum_total spend={t_ps:.2f} leads={t_pl_i} | "
        f"delta spend={100*(t_ps-t_ms)/t_ms:+.3f}% leads={100*(t_pl_i-t_ml_i)/t_ml_i:+.3f}%"
    )


if __name__ == "__main__":
    main()
